#!/usr/bin/env python3
"""Independent source validation and forensic review of Swedish+French gaps.

Downloads real government sources and validates exact food and source nutrient
semantics before writing evidence. No runtime or live product-file changes.
"""
import argparse,hashlib,io,json,math,re,requests
from pathlib import Path
from openpyxl import load_workbook

def get(url,accept="*/*"):
 r=requests.get(url,timeout=60,headers={"Accept":accept});r.raise_for_status();return r
def observed(old,record,field,number):
 assert record["nutrients"][field] is None,field
 assert record["primary_source"]["source_record_id"]==old,(field,old,record["primary_source"])
def macro_match(main,prot,fat):
 a=main["nutrients"]
 return abs(a["protein_per_100g"]-prot)<=max(.5,.2*a["protein_per_100g"]) and abs(a["fat_per_100g"]-fat)<=max(.5,.2*a["fat_per_100g"])
def main():
 ap=argparse.ArgumentParser()
 for k in ("blocked","out"):ap.add_argument("--"+k,required=True)
 a=ap.parse_args()
 foods={p["family_key"]:p for p in json.loads(Path(a.blocked).read_text())}
 out=Path(a.out);out.mkdir(exist_ok=True,parents=True)
 proofs=[];excluded=[]
 def sv(food_id):
  resp=get("https://dataportal.livsmedelsverket.se/livsmedel/api/v1/livsmedel/"+str(food_id)+"/naringsvarden?sprak=2","application/json")
  data=resp.json();assert isinstance(data,list)
  return {x["euroFIRkod"]:x for x in data}
 z=foods["zander"];observed("T603100",z,"selenium_ug",22.6)
 rz=sv(1263);se=rz["SE"];assert se["varde"]==22.6 and se["enhet"]=="µg" and se["matrisenhetkod"]=="W"
 assert se["metodtypkod"]=="A" and se["vardetypkod"]=="MN"
 assert macro_match(z,rz["PROT"]["varde"],rz["FAT"]["varde"])
 proofs.append({"family_key":"zander","field":"selenium_ug","value":22.6,
  "source":"SWEDEN_SLV:1263","food":"Pike-perch raw; Sander lucioperca",
  "source_method":"Independent laboratory, analytical ICP mineral, a composite specimen", 
  "source_value_type_code":"MN", "source_method_code":"A",
  "provenance":{k:se.get(k) for k in ("vardetyp","ursprung","metodtyp","publikation")},
  "quality":"MEASURED_DONOR_SAME_RAW_SPECIES_MACROS_MATCHED"})
 bw=foods["flour buckwheat"];observed("C424000",bw,"vitamin_c_mg",0)
 rb=sv(1930);vc=rb["VITC"];assert vc["varde"]==0 and vc["enhet"]=="mg" and vc["matrisenhetkod"]=="W"
 assert vc["vardetypkod"]=="LZ" and vc["metodtypkod"]=="U"
 assert macro_match(bw,rb["PROT"]["varde"],rb["FAT"]["varde"])
 proofs.append({"family_key":"flour buckwheat","field":"vitamin_c_mg","value":0,
  "source":"SWEDEN_SLV:1930","food":"Buckwheat flour; Fagopyrum esculentum",
  "source_method":"logical zero (not chemically measured)",
  "source_value_type_code":"LZ","source_method_code":"U",
  "quality":"SOURCE_DERIVED_LOGICAL_ZERO_NOT_MEASURED"})
 bse=rb["SE"];assert bse["varde"]==0 and bse["vardetypkod"]=="BL"
 excluded.append({"family_key":"flour buckwheat","field":"selenium_ug",
   "source":"SWEDEN_SLV:1930","reported_value":0,
   "reason":"Below limit of detection (BL), does not justify exact 0 mg or ug",
   "status":"LEFT_UNKNOWN_CENSORED"})
 # French Ciqual 2025, independently checked MD5, exact cited row IDs.
 fr=get("https://entrepot.recherche.data.gouv.fr/api/access/datafile/666260").content
 assert hashlib.md5(fr).hexdigest()=="0d9758ce23f3f13dd63a005bc1bb4f2c"
 sheet=load_workbook(io.BytesIO(fr),read_only=True,data_only=True)["composition nutritionnelle"]
 iterator=sheet.iter_rows(values_only=True);hdr=list(next(iterator))
 table={str(r[6]):dict(zip(hdr,r)) for r in iterator if r[6]}
 def col(q):return next(h for h in hdr if q in re.sub(r"\s+"," ",str(h)).casefold())
 barley=foods["flour barley"];observed("C243000",barley,"selenium_ug",1)
 fb=table["9550"];assert fb["alim_nom_fr"]=="Farine d'orge"
 val=fb[col("sélénium")];assert str(val).replace(",",".")=="1"
 assert macro_match(barley,float(str(fb[col("protéines")]).replace(",",".")),float(str(fb[col("lipides")]).replace(",",".")))
 proofs.append({"family_key":"flour barley","field":"selenium_ug","value":1.0,
  "source":"FRANCE_CIQUAL:9550","food":"Farine d'orge / barley flour",
  "source_method":"Published source value, full per-entry laboratory method not yet traced",
  "source_value_type_code":"SOURCE_REPORTED_WITH_ORIGIN_REVIEW",
  "quality":"MATCHING_FOOD_MACROS_SOURCE_REPORTED_NOT_SAME_LAB"})
 leaf=foods["spice bay leaf"];observed("170917",leaf,"sugar_per_100g",48.6)
 fl=table["11053"];assert fl["alim_nom_fr"]=="Laurier, feuille"
 sugar=float(str(fl[col("sucres")]).replace(",","."))
 assert sugar==48.6
 carbo=leaf["nutrients"]["carbs_per_100g"];fibre=leaf["nutrients"]["fiber_per_100g"]
 excluded.append({"family_key":"spice bay leaf","field":"sugar_per_100g",
     "source":"FRANCE_CIQUAL:11053","reported_value":sugar,
     "reason":"Implausible total sugars: 48.6 g equals near-exactly USDA total carbohydrates 74.97 g less fiber 26.3 g; possible carried-over carbohydrate-by-difference, needs source-method review",
     "USDA_total_carbs":carbo,"USDA_fibre":fibre,
     "status":"QUARANTINED_SUSPECT_CARBOHYDRATE_DOUBLE_COUNT"})
 pepper=foods["spice pepper white"];observed("170933",pepper,"vitamin_k_mcg",0)
 fp=table["11019"];assert fp["alim_nom_fr"]=="Poivre blanc, poudre"
 k1=fp[col("vitamine k1")];assert str(k1) in ("0","0.0")
 excluded.append({"family_key":"spice pepper white","field":"vitamin_k_mcg",
  "source":"FRANCE_CIQUAL:11019","reported_value":0,
  "reason":"Source explicitly says Vitamin K1 (phylloquinone), calculator field may indicate combined vitamin K. Ciqual value origin/method not verified.",
  "status":"HELD_FOR_K1_VERSUS_VITAMIN_K_SEMANTICS_REVIEW"})
 report={
 "status":"SCIENCE_SOURCE_VALIDATED_REVIEW_REQUIRED",
 "new_nutrient_observations":len(proofs),
 "provenance_strict":proofs,"quarantined_data":excluded,
 "original_live_base_unchanged":True,
 "original_blocked_records_unchanged":True}
 (out/"gov_sources_added_fields_check.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
 print("GOVERNMENT_SOURCE_NEW_FIELDS",json.dumps({"count":len(proofs),"fields":proofs,"quarantined":excluded},ensure_ascii=False))
if __name__=="__main__":main()
