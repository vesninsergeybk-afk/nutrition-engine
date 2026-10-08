#!/usr/bin/env python3
"""Rank UK CoFID2021 exact/near food description and per-nutrient observations.

Only extracts publication values as science review suggestions.
Never infers missing values, zeroes, or swaps food processing states.
"""
from __future__ import annotations
import argparse,json,re,unicodedata,math
from pathlib import Path
from difflib import SequenceMatcher
from openpyxl import load_workbook
from collections import Counter,defaultdict

def norm(t):
    t=unicodedata.normalize("NFKC",str(t or "")).casefold()
    t=re.sub(r"\bspices?\b"," ",t)
    t=re.sub(r"\bvegetable oil\b","oil",t)
    t=re.sub(r"\boils?\b","oil",t)
    t=re.sub(r"\bnuts?\b"," ",t)
    t=re.sub(r"\bseeds?\b","seed",t)
    t=re.sub(r"[^a-z0-9]+"," ",t)
    return " ".join(t.split())
def sim(a,b):
    a,b=norm(a),norm(b)
    x,y=set(a.split()),set(b.split())
    return (2*len(x&y)/(len(x)+len(y)) if (x or y) else 0)*.7+SequenceMatcher(None,a,b).ratio()*.3
def suspicious(a,b):
    a,b=set(norm(a).split()),set(norm(b).split())
    contrasts=[("raw","cooked"),("raw","boiled"),("raw","roasted"),
         ("raw","fried"),("raw","dried"),("fresh","dried"),("whole","skimmed"),
         ("low","high"),("sweetened","unsweetened"),("canned","dried"),
         ("canned","raw"),("whole","refined")]
    return [f"{x}/{y}" for x,y in contrasts if x in a and y in b or y in a and x in b]
def parse_number(v):
    if type(v) in (int,float) and math.isfinite(v):return float(v)
    if not isinstance(v,str):return None
    v=v.strip()
    if v.lower() in ("n","tr","trace","-","", "nil"):return None
    try:
        f=float(v)
        return f if math.isfinite(f) else None
    except (ValueError,TypeError):return None
MAP={
  "sugar_per_100g":("1.3 Proximates",["Total sugars (g)"]),
  "fiber_per_100g":("1.3 Proximates",["AOAC fibre (g)","Fibre (g)"]),
  "vitamin_e_mg":("1.5 Vitamins",["Vitamin E (mg)"]),
  "vitamin_k_mcg":("1.5 Vitamins",["Vitamin K1 (µg)"]),
  "vitamin_d_mcg":("1.5 Vitamins",["Vitamin D (µg)"]),
  "vitamin_c_mg":("1.5 Vitamins",["Vitamin C (mg)"]),
  "selenium_ug":("1.4 Inorganics",["Selenium (µg)"]),
  "manganese_mg":("1.4 Inorganics",["Manganese (mg)"]),
  "copper_mg":("1.4 Inorganics",["Copper (mg)"]),
  "vitamin_b5_mg":("1.5 Vitamins",["Pantothenic acid (mg)"]),
  "vitamin_b9_mcg":("1.5 Vitamins",["Total folate (µg)"]),
  "vitamin_b12_mcg":("1.5 Vitamins",["Vitamin B12 (µg)"])
}
def main():
  ap=argparse.ArgumentParser()
  for arg in ("cofid","blocked","out-dir"):ap.add_argument("--"+arg,required=True)
  a=ap.parse_args()
  p=json.loads(Path(a.blocked).read_text())
  wb=load_workbook(a.cofid,read_only=True,data_only=True)
  sample=wb["1.3 Proximates"]
  basic={str(row[0]):{"food_code":row[0],"name":row[1],"description":row[2],
      "references":row[5]} for row in sample.iter_rows(min_row=4,values_only=True) if row[0] is not None and row[1]}
  tabs={}
  for sheet in ("1.3 Proximates","1.4 Inorganics","1.5 Vitamins"):
    ws=wb[sheet]
    head=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
    colnames={re.sub(r"\s+"," ",str(k)).strip().casefold():i for i,k in enumerate(head) if isinstance(k,str)}
    rows={str(row[0]):row for row in ws.iter_rows(min_row=4,values_only=True) if row[0]}
    tabs[sheet]={"colnames":colnames,"rows":rows}
  res=[];candidates=[]
  for x in p:
    source=x["primary_source"]["source_name"]
    top=sorted(basic.values(),key=lambda f:sim(source,f["name"]),reverse=True)[:5]
    rec={"family_key":x["family_key"],"source_name":source,
      "nearby_foods":[]}
    for food in top:
      if sim(source,food["name"])<.44:continue
      name=food["name"]
      exact=norm(source)==norm(name)
      contrasts=suspicious(source,name)
      food_obs=[]
      for nutrient in [k for k,v in x["nutrients"].items() if v is None and k in MAP]:
        sheet,labels=MAP[nutrient]
        row=tabs[sheet]["rows"].get(str(food["food_code"]))
        cols=tabs[sheet]["colnames"]
        if row is None:continue
        for label in labels:
          j=cols.get(label.casefold())
          if j is None:continue
          v=parse_number(row[j])
          if v is None:break
          obs={"field":nutrient,"value":v,"column":label,
             "food_code":food["food_code"],"food_name":name,"food_description":food["description"],
             "food_reference":food["references"],"source":"UK CoFID 2021",
             "id_mode":"NORMALIZED_EXACT" if exact else "REQUIRES_HUMAN_IDENTITY_REVIEW",
             "processing_state_conflicts":contrasts,"status":"REVIEW_ONLY_NO_IMPORT"}
          food_obs.append(obs);candidates.append({"family":x["family_key"],**obs})
          break
      rec["nearby_foods"].append({
         "cofid_id":food["food_code"],"cofid_name":name,
         "similarity":round(sim(source,name),4),
         "normalized_exact":exact,
         "state_conflicts":contrasts,
         "missing_nutrient_observations":food_obs
      })
    res.append(rec)
  report={"status":"REVIEW_ONLY_NO_VALUES_IMPORTED",
    "blocked_foods":len(p),"matching_cofid_names":sum(bool(x["nearby_foods"]) for x in res),
    "potential_source_values":len(candidates),
    "normalized_exact_candidate_observations":sum(x["id_mode"]=="NORMALIZED_EXACT" and not x["processing_state_conflicts"] for x in candidates),
    "value_fields":dict(Counter(x["field"] for x in candidates)),
    "unmapped_fields":["choline_mg","added_sugar","unsat"],
    "source":"Public Health England UK CoFID 2021","foods":res}
  out=Path(a.out_dir);out.mkdir(parents=True,exist_ok=True)
  (out/"uk_cofid_56_semantic_candidates.json").write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str)+"\n")
  print("COFID_56_SUMMARY",json.dumps({k:v for k,v in report.items() if k!="foods"},ensure_ascii=False))
  for f in res:
    if f["nearby_foods"]:
      print("COFID_56_FAMILY",json.dumps(f,ensure_ascii=False,default=str)[:6500])
if __name__=="__main__":main()
