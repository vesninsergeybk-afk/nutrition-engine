#!/usr/bin/env python3
"""Independent verify exact saffron sugar observation from UK CoFID (staging).

The same ingredient must match proximate nutrients before approval. No
blind mapping from unrelated product descriptions.
"""
import argparse,json,math
from pathlib import Path
from openpyxl import load_workbook

def number(v):
  try:return float(v)
  except (ValueError,TypeError):return None

def main():
  ap=argparse.ArgumentParser()
  for x in ("cofid","blocked","output"):ap.add_argument("--"+x,required=True)
  a=ap.parse_args()
  blocked=json.loads(Path(a.blocked).read_text())
  p=next(p for p in blocked if p["family_key"]=="spice saffron")
  assert p["primary_source"]["source_name"]=="Spices, saffron"
  assert str(p["primary_source"]["source_record_id"])=="170934"
  assert p["nutrients"]["sugar_per_100g"] is None
  wb=load_workbook(a.cofid,read_only=True,data_only=True)
  ws=wb["1.3 Proximates"]
  head=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
  indexes={str(k).strip():i for i,k in enumerate(head) if k}
  matching=[row for row in ws.iter_rows(min_row=4,values_only=True) if row[0]=="13-852"]
  assert len(matching)==1
  row=matching[0]
  assert row[1]=="Saffron"
  sugar=number(row[indexes["Total sugars (g)"]])
  assert sugar==42.4
  comp=[]
  pairs={"kcal":"Energy (kcal) (kcal)","protein_per_100g":"Protein (g)","fat_per_100g":"Fat (g)","carbs_per_100g":"Carbohydrate (g)"}
  for field,col in pairs.items():
    v=number(row[indexes[col]])
    old=p["nutrients"][field]
    discrepancy=abs(old-v)/max(1.0,old) if v is not None else None
    comp.append({"field":field,"usda_value":old,"cofid_value":v,"discrepancy_fraction":discrepancy})
  # Product equivalence is not guaranteed by names. Keep as candidate when
  # a major proximate mismatch appears, even if UK database lists exact spice.
  major=[x for x in comp if x["cofid_value"] is None or x["discrepancy_fraction"]>0.20]
  report={
    "status":"NEEDS_SCIENTIFIC_REVIEW" if major else "SAME_FOOD_MACRO_PARITY_REVIEW_REQUIRED",
    "family_key":"spice saffron",
    "source_usda":"USDA_FDC:170934 / Spices, saffron",
    "source_cofid":"UK CoFID 2021 food 13-852 / Saffron",
    "field":"sugar_per_100g",
    "value":42.4,
    "unit":"g/100g",
    "source_column":"Total sugars (g)",
    "cofid_reference":row[5],
    "macro_comparison":comp,
    "large_discrepancies":major,
    "scientific_status":"EVIDENCE_STAGING_ONLY",
    "production_authorized":False
  }
  out=Path(a.output);out.mkdir(exist_ok=True,parents=True)
  (out/"saffron_sugar42_4_cofid_parity.json").write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str)+"\n")
  print("SAFFRON_SUGAR_COFID",json.dumps(report,ensure_ascii=False,default=str))
if __name__=="__main__":main()
