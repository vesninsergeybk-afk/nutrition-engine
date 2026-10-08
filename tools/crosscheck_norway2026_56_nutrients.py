#!/usr/bin/env python3
"""2026 Norwegian government table, 56 missing field crosswalk and provenance.

This is a source study. Reports values with per-nutrient source code. The file
is never used to rewrite runtime. Direct same food identities and macros
required before any values move from candidates to independent review.
"""
from __future__ import annotations
import argparse,io,json,math,re,requests,hashlib,unicodedata
from pathlib import Path
from collections import Counter
from difflib import SequenceMatcher
from openpyxl import load_workbook
URL="https://mattilsynet-xp7prod.enonic.cloud/en/_/attachment/inline/9b62762f-ccf5-4194-bfd0-c375b9d2ddc8:33207f2fabe06b025f3b04110d4a8f20b077f23c/Norwegian%20Food%20Composition%20Table%202026.xlsx"
def norm(s):
 s=unicodedata.normalize("NFKD",str(s or "").lower())
 s="".join(c for c in s if not unicodedata.combining(c))
 s=s.replace("spices","").replace("spice","").replace("seeds","seed")
 return " ".join(re.sub(r"[^a-z0-9]+"," ",s).split())
def score(a,b):
 a=norm(a);b=norm(b);x=set(a.split());y=set(b.split())
 if not x|y:return 0
 return round(.7*len(x&y)/len(x|y)+.3*SequenceMatcher(None,a,b).ratio(),4)
def exact(a,b):return sorted(norm(a).split())==sorted(norm(b).split())
def state(a,b):
 a=set(norm(a).split());b=set(norm(b).split())
 return [p+"/"+q for p,q in (("raw","cooked"),("fresh","dried"),("raw","dried"),
  ("whole","skimmed"),("canned","raw"),("dry","cooked"),("fresh","canned"),
  ("sweetened","unsweetened"),("oil","seed"),("canned","dry"),("high","low"))
  if (p in a and q in b) or (q in a and p)]
def val(v):
 if isinstance(v,(float,int)) and math.isfinite(v) and v>=0:return float(v)
 if not isinstance(v,str) or not v:return None
 v=v.strip().replace(",",".")
 if v.startswith(("<",">","~")):return None
 try:
  a=float(v);return a if math.isfinite(a) and a>=0 else None
 except ValueError:return None
def read_sheet(book,name):
 ws=book[name]
 it=ws.iter_rows(values_only=True)
 header=list(next(it))
 return header,{str(r[0]):dict(zip(header,r)) for r in it if r[0] and re.match(r"^\d+\.\d+$",str(r[0]))}
def main():
 p=argparse.ArgumentParser()
 for opt in ("blocked","out"):p.add_argument("--"+opt,required=True)
 a=p.parse_args()
 blocked=json.loads(Path(a.blocked).read_text())
 raw=requests.get(URL,timeout=75);raw.raise_for_status()
 assert raw.content[:2]==b"PK"
 digest=hashlib.sha256(raw.content).hexdigest()
 book=load_workbook(io.BytesIO(raw.content),read_only=True,data_only=True)
 head,foods=read_sheet(book,"Foods (all nutrients)")
 sh,source=read_sheet(book,"Sources (all nutrients)")
 lookup={str(row[0]):row[1] for row in book["Source Lookup"].iter_rows(values_only=True) if row[0]}
 assert len(foods)==2121,len(foods)
 print("NORWAY2026_ALL_NUTRIENT_COLUMNS",json.dumps(head,ensure_ascii=False))
 print("NORWAY2026_SOURCETYPE_CODES",json.dumps({k:lookup[k] for k in list(lookup)[:25]},ensure_ascii=False))
 cands={
 "choline_mg":[r"^choline"],
 "vitamin_k_mcg":[r"^vitamin k"],
 "vitamin_e_mg":[r"^vitamin e"],
 "vitamin_b5_mg":[r"^pantothenic",r"^vitamin b5"],
 "selenium_ug":[r"^selenium"],
 "sugar_per_100g":[r"^sugar"],
 "vitamin_b9_mcg":[r"^folate"],
 "vitamin_b12_mcg":[r"^vitamin b12",r"^vitamin b-12"],
 "vitamin_c_mg":[r"^vitamin c"],
 "vitamin_d_mcg":[r"^vitamin d"],
 "fiber_per_100g":[r"^dietary fibre"],
 "manganese_mg":[r"^manganese"],
 "copper_mg":[r"^copper"],
 "vitamin_a_mcg":[r"^vitamin a"],
 "vitamin_b2_mg":[r"^riboflavin"],
 "sfa":[r"^saturated fatty acids"],
 "unsat":[r"^monounsaturated fatty acids",r"^polyunsaturated fatty acids"]
 }
 fieldcols={k:[h for h in head if any(re.search(pattern,norm(h)) for pattern in ps)] for k,ps in cands.items()}
 print("NORWAY2026_FIELD_MAP",json.dumps(fieldcols,ensure_ascii=False))
 macro_cols={"protein_per_100g":"Protein (g)","fat_per_100g":"Fat (g)"}
 ids=list(foods)
 indexed=[];counts=Counter();byfield=Counter();observations=[]
 for t in blocked:
  src=t["primary_source"]["source_name"]
  matches=sorted(ids,key=lambda k:score(src,foods[k]["Matvare"]),reverse=True)[:6]
  found=[]
  for k in matches:
   f=foods[k];name=f["Matvare"];sim=score(src,name);name_equal=exact(src,name)
   if sim<.48:continue
   macro=[]
   for field,col in macro_cols.items():
    old=t["nutrients"][field];v=val(f[col])
    macro.append({"field":field,"old":old,"new":v,
     "consistent":v is not None and abs(old-v)<=max(.5,.20*old)})
   mp=all(x["consistent"] for x in macro)
   issues=state(src,name)
   entry={"norwegian_id":k,"norwegian_name":name,"name_score":sim,
     "canonical_name_exact":name_equal,"macro_comparison":macro,"macro_parity":mp,
     "processing_conflicts":issues,"missing_fields":[]}
   for field,orig in t["nutrients"].items():
    if orig is not None:continue
    columns=fieldcols.get(field,[])
    if len(columns)!=1:continue
    col=columns[0];v=val(f[col]);code=str(source[k].get(col,""))
    if v is None or not code or code=="10":continue
    record={"family":t["family_key"],"field":field,"value":v,"food_id":k,"food_name":name,
      "source_column":col,"source_code":code,"source_explanation":lookup.get(code),
      "food_name_exact":name_equal,"macro_parity":mp,
      "processing_conflicts":issues,"status":"REQUIRES_SCIENTIFIC_REVIEW"}
    # 60b: published as below quantification, do not promote to numeric zero
    if code in ("60a","60b","60c"):record["status"]="CENSORED_DO_NOT_INTERPRET_AS_EXACT"
    elif name_equal and mp and not issues:
      record["status"]="SAME_FOOD_MACRO_MATCHED_SOURCE_REPORTED"
      counts[t["family_key"]]+=1;byfield[field]+=1
      observations.append(record)
    entry["missing_fields"].append(record)
   found.append(entry)
  indexed.append({"family":t["family_key"],"reference_food":src,"matches":found})
  if found:
    print("NORWAY2026_FOOD",json.dumps({"family":t["family_key"],"reference_food":src,
      "top":[{"name":m["norwegian_name"],"code":m["norwegian_id"],"exact":m["canonical_name_exact"],
      "macro":m["macro_parity"],"candidate_fields":len(m["missing_fields"])} for m in found[:5]]},ensure_ascii=False))
 report={"source":"Norwegian Food Composition Table 2026, Mattilsynet",
  "source_url":URL,"sha256":digest,"food_count":len(foods),
  "source_method_codes":lookup,"column_map":fieldcols,
  "strict_nutrient_observations":len(observations),
  "strict_target_families":len(counts),"strict_by_field":dict(byfield),
  "observations":observations,"target_food_candidates":indexed,
  "runtime_changes":0,"staging_changes":0}
 out=Path(a.out);out.mkdir(exist_ok=True,parents=True)
 (out/"norway2026_56_crosswalk.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
 print("NORWAY2026_CROSSWALK_SUMMARY",json.dumps({k:v for k,v in report.items() if k not in ("source_method_codes","column_map","observations","target_food_candidates")},ensure_ascii=False))
 for x in observations:print("NORWAY2026_SOURCE_NUTRIENT",json.dumps(x,ensure_ascii=False))
if __name__=="__main__":main()
