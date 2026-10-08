#!/usr/bin/env python3
"""Inspect national Scandinavian and French nutrient sources, research only.

CI downloads primary source lists, candidate records and nutrient definitions.
No updates to runtime or staging numeric fields in this first source-inspection.
"""
import argparse,json,re,unicodedata,requests,hashlib,io
from pathlib import Path
from openpyxl import load_workbook
from difflib import SequenceMatcher

def norm(x):
  return " ".join(re.sub(r"[^a-z0-9]+"," ",unicodedata.normalize("NFKC",str(x or "")).lower()).split())
def score(a,b):
  a=norm(a);b=norm(b);x=set(a.split());y=set(b.split())
  return round(.65*len(x&y)/len(x|y)+.35*SequenceMatcher(None,a,b).ratio(),4) if x|y else 0
def loadurl(url):
  r=requests.get(url,timeout=45,headers={"User-Agent":"nutrition-engine-research/1.0 (academic source quality audit)","Accept":"application/json"})
  print("HTTP",r.status_code,url,"bytes",len(r.content))
  r.raise_for_status()
  return r
def main():
  p=argparse.ArgumentParser()
  for x in ("blocked","out"):p.add_argument("--"+x,required=True)
  a=p.parse_args()
  dst=Path(a.out);dst.mkdir(parents=True,exist_ok=True)
  blocked=json.loads(Path(a.blocked).read_text())
  report={"status":"INSPECTION_NOT_APPROVAL","blocked_food_count":len(blocked),
          "source_schema":{},"nutrient_fields_to_search":{}}
  for b in blocked:
    report["nutrient_fields_to_search"][b["family_key"]]={
      "name":b["primary_source"]["source_name"],
      "missing":[f for f,v in b["nutrients"].items() if v is None],
      "source":b["primary_source"]["source_registry_id"]+":"+str(b["primary_source"]["source_record_id"])}
  (dst/"blocked56_missing_field_index.json").write_text(json.dumps(report["nutrient_fields_to_search"],ensure_ascii=False,indent=2)+"\n")
  # Sweden agency open CC-BY API, English naming
  try:
    sw=loadurl("https://dataportal.livsmedelsverket.se/livsmedel/api/v1/livsmedel?offset=0&limit=2500&sprak=2").json()
    report["source_schema"]["sweden"]={
      "json_type":type(sw).__name__,
      "keys":list(sw)[:30] if isinstance(sw,dict) else None,
      "sample":str(sw)[:2500]}
    (dst/"sweden_api_list_preview.json").write_text(json.dumps(sw,ensure_ascii=False,indent=2)[:220000]+"\n")
    print("SWEDEN_LIST_STRUCTURE",json.dumps(report["source_schema"]["sweden"],ensure_ascii=False)[:2500])
  except Exception as e:
    report["source_schema"]["sweden"]={"error":type(e).__name__+": "+str(e)[:250]}
    print("SWEDEN_SOURCE_ERROR",report["source_schema"]["sweden"]["error"])
  # France official Ciqual2025 data via standard Dataverse metadata endpoint
  try:
    api="https://entrepot.recherche.data.gouv.fr/api/datasets/:persistentId/?persistentId=doi%3A10.57745%2FRDMHWY"
    fr=loadurl(api).json()
    filelist=[{"id":f.get("dataFile",{}).get("id"),"filename":f.get("dataFile",{}).get("filename"),
              "md5":f.get("dataFile",{}).get("md5")} for f in fr["data"]["latestVersion"]["files"]]
    report["source_schema"]["ciqual"]={"files":filelist}
    print("CIQUAL_FILE_INDEX",json.dumps(filelist,ensure_ascii=False))
    selected=next(f for f in filelist if f["filename"]=="Table Ciqual 2025_FR_2025_11_03.xlsx")
    data=loadurl("https://entrepot.recherche.data.gouv.fr/api/access/datafile/"+str(selected["id"])).content
    assert hashlib.md5(data).hexdigest()==selected["md5"],"Ciqual MD5 mismatch"
    w=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    tabs=[]
    for sheet in w:
      top=[[str(x)[:90] for x in row[:12]] for row in sheet.iter_rows(min_row=1,max_row=min(7,sheet.max_row),values_only=True)]
      tabs.append({"sheet":sheet.title,"rows":sheet.max_row,"columns":sheet.max_column,"top7":top})
    report["source_schema"]["ciqual"]["tabs"]=tabs
    print("CIQUAL2025_SHEETS",json.dumps(tabs,ensure_ascii=False)[:9000])
  except Exception as e:
    report["source_schema"]["ciqual"]={"error":type(e).__name__+": "+str(e)[:350]}
    print("CIQUAL_SOURCE_ERROR",report["source_schema"]["ciqual"]["error"])
  (dst/"national_2025_2026_sources_schema.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n")
if __name__=="__main__":main()
