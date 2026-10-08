#!/usr/bin/env python3
"""Verify AFCD Release 3 arrowroot selenium zero (published) against official XLSX.

A published zero may be rounded/below limit. Does not establish absence.
"""
import argparse,json,re
from pathlib import Path
from openpyxl import load_workbook
def records(p,sheet):
  ws=load_workbook(p,read_only=True,data_only=True)[sheet]
  head=[re.sub(r"\s+"," ",str(x or "")).strip() for x in next(ws.iter_rows(min_row=3,max_row=3,values_only=True))]
  return {r[0]:dict(zip(head,r)) for r in ws.iter_rows(min_row=4,values_only=True) if r[0]}
def main():
  p=argparse.ArgumentParser()
  for x in ("foods","nutrients","blocked","out"):p.add_argument("--"+x,required=True)
  a=p.parse_args()
  base=json.loads(Path(a.blocked).read_text())
  target=next(x for x in base if x["family_key"]=="flour arrowroot")
  assert target["nutrients"]["selenium_ug"] is None
  assert target["primary_source"]["source_registry_id"]=="GERMANY_BLS"
  assert target["primary_source"]["source_record_id"]=="K550000"
  food=records(a.foods,"Food details")["F003983"]
  nutr=records(a.nutrients,"All solids & liquids per 100 g")["F003983"]
  assert food["Food Name"]=="Flour, arrowroot" and food["Derivation"]=="Analysed"
  assert nutr["Food Name"]=="Flour, arrowroot"
  sel=nutr["Selenium (Se) (ug)"]
  assert isinstance(sel,(float,int)) and sel==0, sel
  assert nutr["Protein (g)"]==0.4 and nutr["Fat, total (g)"]==0.2
  # Absolute macro tolerances require checking tiny amounts explicitly.
  assert abs(target["nutrients"]["protein_per_100g"]-0.4)<0.11
  assert abs(target["nutrients"]["fat_per_100g"]-0.2)<0.11
  out=Path(a.out);out.mkdir(parents=True,exist_ok=True)
  result={"status":"PASS_SOURCE_REPORTED_ZERO_SCIENCE_REVIEW_REQUIRED",
    "family_key":"flour arrowroot","primary_source":"GERMANY_BLS:K550000",
    "donor_source":"AFCD Release 3 F003983","field":"selenium_ug",
    "value":0.0,"unit":"µg/100g",
    "donor_derivation":food["Derivation"],
    "sampling_details":food["Sampling Details"],
    "zero_meaning":"published reportable zero, not definitive evidence of absence",
    "scientific_review_required":True,"production_unchanged":True}
  (out/"arrowroot_selenium_zero_evidence.json").write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n")
  print("ARROWROOT_AFCD_SCIENCE_GATE",json.dumps(result,ensure_ascii=False))
if __name__=="__main__":main()
