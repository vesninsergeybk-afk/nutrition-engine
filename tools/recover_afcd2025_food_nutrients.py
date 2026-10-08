#!/usr/bin/env python3
"""Food-specific AFCD Release 3 nutrient donors, with strict provenance guards.

Creates research-only candidate records for 56 blocked source profiles. Every
candidate references exact printed AFCD column, unit, and official Food Key.
Strict candidates require same-food identity or an explicit source FDC ID in
AFCD's sampling details, macro parity, and appropriate processing state.
"""
import argparse,json,re,math,unicodedata
from collections import Counter
from pathlib import Path
from difflib import SequenceMatcher
from openpyxl import load_workbook

def norm(s):
    s=unicodedata.normalize("NFKC",str(s or "")).casefold()
    s=s.replace("coconut meat","coconut").replace("catfish (wolffish)","wolffish")
    s=re.sub(r"\bspices?\b"," ",s)
    s=re.sub(r"\bvegetable\b"," ",s)
    s=re.sub(r"\boils?\b","oil",s)
    s=re.sub(r"\bseeds?\b","seed",s)
    s=re.sub(r"\bflours?\b","flour",s)
    s=re.sub(r"[^a-z0-9]+"," ",s)
    return " ".join(s.split())
def similarity(a,b):
    aa,bb=norm(a),norm(b)
    a,b=set(aa.split()),set(bb.split())
    return round(.6*(len(a&b)/len(a|b) if a|b else 0)+.4*SequenceMatcher(None,aa,bb).ratio(),4)
def name_exact(a,b):
    return sorted(norm(a).split())==sorted(norm(b).split())
def state_conflicts(a,b):
    a=set(norm(a).split());b=set(norm(b).split())
    pairs=(("raw","cooked"),("raw","boiled"),("raw","roasted"),
        ("raw","dried"),("fresh","dried"),("dry","cooked"),("whole","skimmed"),
        ("seed","oil"),("toasted","raw"),("canned","raw"),("flaked","creamed"),
        ("high","low"),("sweetened","unsweetened"),("canned","dry"))
    return [f"{x}_vs_{y}" for x,y in pairs if (x in a and y in b) or (y in a and x in b)]
def numeric(v):
    return float(v) if isinstance(v,(int,float)) and math.isfinite(v) and v>=0 else None
def rows(book,sheet,start=4):
    ws=book[sheet]
    heads=[re.sub(r"\s+"," ",str(x or "")).strip() for x in next(ws.iter_rows(min_row=3,max_row=3,values_only=True))]
    return heads,[dict(zip(heads,vals)) for vals in ws.iter_rows(min_row=start,values_only=True) if vals[0]]
def fieldmap(headers):
    field_patterns={
      "sugar_per_100g":[r"^total sugars \(g\)$"],
      "added_sugar":[r"^added sugars \(g\)$"],
      "fiber_per_100g":[r"^total dietary fibre \(g\)$"],
      "selenium_ug":[r"^selenium \(se\).*"],
      "choline_mg":[r"^choline.*\(mg\)$"],
      "vitamin_e_mg":[r"^vitamin e.*\(mg\)$",r"^alpha.tocopherol.*\(mg\)$"],
      "vitamin_k_mcg":[r"^vitamin k.*\(.*g\)$"],
      "vitamin_b5_mg":[r"^pantothenic acid.*\(mg\)$"],
      "vitamin_b9_mcg":[r"^folate.*\(.*g\)$"],
      "vitamin_c_mg":[r"^vitamin c.*\(mg\)$"],
      "vitamin_a_mcg":[r"^vitamin a.*\(.*g\)$",r"^retinol equivalents.*\(.*g\)$"],
      "vitamin_d_mcg":[r"^vitamin d.*\(.*g\)$"],
      "vitamin_b12_mcg":[r"^vitamin b12.*\(.*g\)$"],
      "manganese_mg":[r"^manganese.*\(mg\)$"],
      "copper_mg":[r"^copper.*\(mg\)$"],
      "vitamin_b2_mg":[r"^riboflavin.*\(mg\)$"],
      "unsat":[r"^fatty acids,.*unsaturated.*\(g\)$"],
    }
    found={}
    for field,patterns in field_patterns.items():
        found[field]=[h for h in headers if any(re.search(p,h,re.IGNORECASE) for p in patterns)]
    return found

def main():
    p=argparse.ArgumentParser()
    for x in ("foods","nutrients","blocked","out"):p.add_argument("--"+x,required=True)
    x=p.parse_args()
    info=json.loads(Path(x.blocked).read_text())
    fbook=load_workbook(x.foods,read_only=True,data_only=True)
    nbook=load_workbook(x.nutrients,read_only=True,data_only=True)
    fhead,foods=rows(fbook,"Food details")
    nhead,nuts=rows(nbook,"All solids & liquids per 100 g")
    nf={p["Public Food Key"]:p for p in nuts}
    assert len(foods)>=1588
    macrofields={"protein_per_100g":"Protein (g)","fat_per_100g":"Fat, total (g)"}
    ncols=fieldmap(nhead)
    print("AFCD_COLUMN_MAP",json.dumps(ncols,ensure_ascii=False))
    print("AFCD_ALL_NUTRIENT_HEADINGS",json.dumps(nhead[:115],ensure_ascii=False))
    results=[];possible=[];strict=[]
    for target in info:
        src=target["primary_source"];name=src["source_name"]
        fdc=str(src["source_record_id"]) if src["source_registry_id"]=="USDA_FDC" else None
        sortedfood=sorted(foods,key=lambda r: similarity(name,r["Food Name"]),reverse=True)
        same_fdc=[f for f in foods if fdc and re.search(r"(?<!\d)"+re.escape(fdc)+r"(?!\d)",str(f.get("Sampling Details") or ""))]
        ids={f["Public Food Key"] for f in sortedfood[:7]}
        ids.update(f["Public Food Key"] for f in same_fdc)
        candidates=[]
        for food in foods:
            if food["Public Food Key"] not in ids:continue
            key=food["Public Food Key"];n=nf.get(key)
            if n is None:continue
            s=similarity(name,food["Food Name"])
            exact=name_exact(name,food["Food Name"])
            same_id=bool(fdc and re.search(r"(?<!\d)"+re.escape(fdc)+r"(?!\d)",str(food.get("Sampling Details") or "")))
            issues=state_conflicts(name,food["Food Name"])
            macro=[]
            for dest,h in macrofields.items():
                val=numeric(n.get(h));reference=target["nutrients"].get(dest)
                if val is not None and reference is not None:
                    macro.append({"field":dest,"main":reference,"afcd":val,
                        "difference_pct":round(100*abs(val-reference)/max(reference,0.1),2),
                        "within_20pct":abs(val-reference)<=max(0.5,0.2*reference)})
            parity=len(macro)==2 and all(v["within_20pct"] for v in macro)
            entry={"family_key":target["family_key"],"primary_source":src,
                "afcd_key":key,"afcd_food_name":food["Food Name"],
                "derivation":food.get("Derivation"),
                "sampling_details":str(food.get("Sampling Details") or "")[:1500],
                "similarity":s,"normalized_name_exact":exact,
                "source_fdc_id_evidenced":same_id,
                "state_conflicts":issues,"macro_check":macro,
                "macro_parity":parity,
                "missing_fields_with_values":[]}
            for field in [k for k,v in target["nutrients"].items() if v is None]:
                columns=ncols.get(field,[])
                if len(columns)!=1:continue
                h=columns[0];val=numeric(n.get(h))
                if val is None:continue
                observation={"field":field,"value":val,"unit_header":h,
                    "source":"Australian Food Composition Database Release 3 Dec 2025",
                    "afcd_food_key":key,"name":food["Food Name"],
                    "name_exact":exact,"shared_fdc_id":same_id,
                    "state_conflicts":issues,
                    "macro_parity":parity,
                    "classification":"REVIEW_ONLY_DO_NOT_IMPORT"}
                entry["missing_fields_with_values"].append(observation)
                possible.append(observation)
                if (exact or same_id) and parity and not issues:
                    strict.append(observation)
            candidates.append(entry)
        candidates.sort(key=lambda y:(y["source_fdc_id_evidenced"],y["normalized_name_exact"],y["macro_parity"],y["similarity"]),reverse=True)
        results.append({"family_key":target["family_key"],"source_name":name,"matches":candidates})
    result={
      "status":"AFCD_2025_CROSSWALK_SCIENTIFIC_REVIEW_ONLY",
      "source_url":"https://www.foodstandards.gov.au/science-data/food-nutrient-databases/afcd/data-files",
      "source_foods":len(foods),"source_nutrient_profiles":len(nuts),
      "blocked_foods":len(info),"source_columns_for_missing_fields":ncols,
      "candidate_numeric_observations":len(possible),
      "strict_identity_and_macro_candidate_observations":len(strict),
      "strict_by_field":dict(Counter(v["field"] for v in strict)),
      "strict_observations":strict,"foods":results,"runtime_changes":0}
    out=Path(x.out);out.mkdir(parents=True,exist_ok=True)
    (out/"afcd2025_missing_nutrient_candidates.json").write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print("AFCD_RECOVERY_SUMMARY",json.dumps({k:v for k,v in result.items() if k not in ("foods","strict_observations","source_columns_for_missing_fields")},ensure_ascii=False))
    for row in strict:
        print("AFCD_STRICT_VALUE",json.dumps(row,ensure_ascii=False))
    for row in results:
        mm=[x for x in row["matches"] if x["normalized_name_exact"] or x["source_fdc_id_evidenced"]]
        if mm:
            print("AFCD_STRICT_FOOD",json.dumps({"family":row["family_key"],"matches":mm},ensure_ascii=False)[:9000])

if __name__=="__main__":main()
